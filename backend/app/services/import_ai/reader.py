"""Turns an arbitrary spreadsheet or CSV file into a compact text grid an AI
can read — the "different file types and formats" input side of AI-assisted
import. Deliberately separate from services/import_excel/, which stays the
fast, deterministic, no-AI path for the one known "Raw data" layout.
"""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl

SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".csv"}

# Keep the prompt payload bounded: enough rows/columns for a multi-year
# monthly-balance sheet, capped overall so a huge or malformed file can't
# blow up token usage or cost.
MAX_ROWS = 400
MAX_COLS = 80
MAX_CHARS = 40_000


class UnsupportedFileTypeError(ValueError):
    pass


def read_file_as_text(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        text = _read_excel_as_text(path)
    elif suffix == ".csv":
        text = _read_csv_as_text(path)
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type {suffix!r}. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n...(truncated)"
    return text


def _read_excel_as_text(path: Path) -> str:
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"=== Sheet: {sheet_name} ===")
        for row in ws.iter_rows(
            min_row=1, max_row=min(ws.max_row, MAX_ROWS), max_col=min(ws.max_column, MAX_COLS)
        ):
            values = ["" if c.value is None else str(c.value) for c in row]
            if any(values):
                parts.append("\t".join(values))
    return "\n".join(parts)


def _read_csv_as_text(path: Path) -> str:
    parts = []
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for i, row in enumerate(reader):
            if i >= MAX_ROWS:
                break
            parts.append("\t".join(row[:MAX_COLS]))
    return "\n".join(parts)
