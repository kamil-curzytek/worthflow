import datetime as dt
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.services import accounts_service, snapshots_service
from app.services.import_excel.importer import commit_import, preview_import
from app.services.import_excel.parser import parse_raw_data_sheet


def build_workbook(
    path: Path, *, include_unknown_row: bool = False, swap_month_columns: bool = False
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Raw data"

    # Year row (3) + month row (4), starting at column D like the real sheet.
    ws["D3"] = 2024
    ws["D4"] = "January"
    ws["E4"] = "February"
    ws["F4"] = "March"

    if swap_month_columns:
        # Reasonable real-world change: columns reordered/labels not in the
        # exact order the importer first saw them in.
        ws["D4"], ws["E4"] = ws["E4"].value, ws["D4"].value

    ws["B5"] = "mBank PLN account"
    ws["D5"] = 1000
    ws["E5"] = 1100
    ws["F5"] = 1200

    ws["B6"] = "N26"
    ws["D6"] = 500
    ws["E6"] = ""  # not tracked that month — must NOT become a zero snapshot
    ws["F6"] = 600

    ws["B7"] = "SUM"  # known derived rollup — must be skipped, not imported
    ws["D7"] = 1500

    if include_unknown_row:
        ws["B8"] = "Some Brand New Account"
        ws["D8"] = 999

    ws["B9"] = "eTrade RSU"
    ws["D9"] = "not-a-number"  # invalid value — must warn, not crash

    wb.save(path)


@pytest.fixture()
def workbook_path(tmp_path):
    path = tmp_path / "test_moneytalks.xlsx"
    build_workbook(path)
    return path


def test_parser_extracts_expected_rows(workbook_path):
    result = parse_raw_data_sheet(str(workbook_path))
    labels = {
        (r.account_mapping.account_name, r.snapshot_date, str(r.value)) for r in result.rows
    }

    assert ("mBank PLN account", dt.date(2024, 1, 1), "1000") in labels
    assert ("mBank PLN account", dt.date(2024, 3, 1), "1200") in labels
    assert ("N26", dt.date(2024, 1, 1), "500") in labels
    # blank cell must not appear at all (not as a zero)
    assert not any(
        r.account_mapping.account_name == "N26" and r.snapshot_date == dt.date(2024, 2, 1)
        for r in result.rows
    )


def test_derived_rollup_row_is_skipped_silently(workbook_path):
    result = parse_raw_data_sheet(str(workbook_path))
    assert not any(r.account_mapping.account_name == "SUM" for r in result.rows)
    assert not any("SUM" in w for w in result.warnings)


def test_invalid_numeric_value_produces_warning_not_crash(workbook_path):
    result = parse_raw_data_sheet(str(workbook_path))
    assert any("non-numeric" in w for w in result.warnings)
    assert not any(r.account_mapping.account_name == "eTrade RSU" for r in result.rows)


def test_unmapped_account_row_produces_warning_not_crash(tmp_path):
    path = tmp_path / "unknown.xlsx"
    build_workbook(path, include_unknown_row=True)
    result = parse_raw_data_sheet(str(path))
    assert any("Some Brand New Account" in w for w in result.warnings)
    assert not any(
        r.account_mapping.account_name == "Some Brand New Account" for r in result.rows
    )


def test_reordered_month_columns_still_parse_correctly(tmp_path):
    path = tmp_path / "reordered.xlsx"
    build_workbook(path, swap_month_columns=True)
    result = parse_raw_data_sheet(str(path))
    # Column D is now labeled February, column E is January — values must
    # follow their own column, not a hardcoded position.
    by_date = {
        r.snapshot_date: r.value
        for r in result.rows
        if r.account_mapping.account_name == "mBank PLN account"
    }
    assert by_date[dt.date(2024, 2, 1)] == Decimal(1000)  # was under D
    assert by_date[dt.date(2024, 1, 1)] == Decimal(1100)  # was under E


def test_preview_import_reports_new_accounts_and_snapshots(db, user, workbook_path):
    preview = preview_import(db, user.id, str(workbook_path))
    assert preview.new_snapshot_count == 5  # 3 mBank + 2 N26 (Feb blank excluded)
    assert preview.duplicate_count == 0
    assert set(preview.new_account_names) == {"mBank PLN account", "N26"}
    assert any("non-numeric" in w for w in preview.warnings)


def test_commit_import_creates_accounts_and_snapshots(db, user, workbook_path):
    batch = commit_import(db, user.id, str(workbook_path))
    assert batch.row_count == 5
    # ImportStatus is a str-enum, so this holds whether SQLAlchemy hands back
    # the enum member or the raw string.
    assert batch.status == "committed"

    accounts = {a.name for a in accounts_service.list_accounts(db)}
    assert accounts == {"mBank PLN account", "N26"}

    mbank = accounts_service.get_account_by_name(db, user.id, "mBank PLN account")
    snaps = snapshots_service.list_snapshots(db, account_id=mbank.id)
    assert len(snaps) == 3


def test_reimporting_same_file_does_not_duplicate(db, user, workbook_path):
    commit_import(db, user.id, str(workbook_path))
    preview = preview_import(db, user.id, str(workbook_path))
    assert preview.new_snapshot_count == 0
    assert preview.duplicate_count == 5

    batch2 = commit_import(db, user.id, str(workbook_path))
    assert batch2.row_count == 0  # nothing new written

    mbank = accounts_service.get_account_by_name(db, user.id, "mBank PLN account")
    snaps = snapshots_service.list_snapshots(db, account_id=mbank.id)
    assert len(snaps) == 3  # still 3, not 6


def test_commit_import_with_overwrite_updates_duplicates(db, user, tmp_path):
    path = tmp_path / "v1.xlsx"
    build_workbook(path)
    commit_import(db, user.id, str(path))

    # "Re-import" with a changed value for the same account/date.
    wb = openpyxl.load_workbook(path)
    wb["Raw data"]["D5"] = 9999
    wb.save(path)

    commit_import(db, user.id, str(path), overwrite_duplicates=True)
    mbank = accounts_service.get_account_by_name(db, user.id, "mBank PLN account")
    jan_snapshot = snapshots_service.find_snapshot(db, mbank.id, dt.date(2024, 1, 1))
    assert jan_snapshot.value == Decimal(9999)


def test_missing_sheet_produces_warning_not_crash(tmp_path):
    path = tmp_path / "wrong_sheet.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Not Raw Data"
    wb.save(path)

    result = parse_raw_data_sheet(str(path))
    assert result.rows == []
    assert any("not found" in w for w in result.warnings)
