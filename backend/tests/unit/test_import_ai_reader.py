import csv

import openpyxl
import pytest

from app.services.import_ai.reader import UnsupportedFileTypeError, read_file_as_text


def test_reads_xlsx_with_sheet_headers(tmp_path):
    path = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Account"
    ws["B1"] = "Jan"
    ws["A2"] = "Revolut"
    ws["B2"] = 1000
    wb.save(path)

    text = read_file_as_text(str(path))
    assert "=== Sheet: Data ===" in text
    assert "Account\tJan" in text
    assert "Revolut\t1000" in text


def test_reads_csv(tmp_path):
    path = tmp_path / "test.csv"
    with path.open("w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Account", "Jan"])
        writer.writerow(["Revolut", "1000"])

    text = read_file_as_text(str(path))
    assert "Account\tJan" in text
    assert "Revolut\t1000" in text


def test_unsupported_extension_raises(tmp_path):
    path = tmp_path / "test.txt"
    path.write_text("hello")
    with pytest.raises(UnsupportedFileTypeError):
        read_file_as_text(str(path))
